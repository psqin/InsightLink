"""生成下载用的静态文档和表格，显式防护 HTML/Markdown 及电子表格公式注入。"""

import csv
import html
import json
import math
import re


FORMATS = {"md", "html", "txt", "json", "csv", "xlsx"}


def text(value, name, maximum, *, empty=False):
    """限制文本的类型、长度和空值；不悄悄截断用户内容。"""
    if not isinstance(value, str) or len(value) > maximum or (not empty and not value.strip()):
        raise ValueError(f"{name}需要{'0' if empty else '1'}–{maximum}字符文本")
    return value


def format_name(value):
    """支持常用别名，同时只允许固定输出格式。"""
    value = {"markdown": "md", "text": "txt"}.get(value, value)
    if value not in FORMATS:
        raise ValueError("文档格式仅支持 md/html/txt/json/csv/xlsx")
    return value


def markdown_literal(value):
    """把外部内容作为字面文本，而不是可执行 HTML、链接或图片语法。"""
    value = html.escape(value, quote=True)
    return re.sub(r"([\\`*_{}\[\]()#+\-.!|>])", r"\\\1", value)


def spreadsheet_literal(value):
    """CSV/Excel 中危险前缀的字符串加单引号，数字本身仍保持数字类型。"""
    if isinstance(value, str) and value.lstrip(" \t\r\n").startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def validate_table(columns, rows):
    """限制表格为 20 列、500 行及标量单元格，拒绝重复列名和非有限数字。"""
    if not isinstance(columns, list) or not 1 <= len(columns) <= 20:
        raise ValueError("表格需要 1–20 列")
    columns = [text(value, "列名", 100) for value in columns]
    if len(set(columns)) != len(columns):
        raise ValueError("表格列名不能重复")
    if not isinstance(rows, list) or len(rows) > 500:
        raise ValueError("表格最多 500 行")
    for row in rows:
        if not isinstance(row, list) or len(row) != len(columns):
            raise ValueError("每行单元格数量必须与列数一致")
        for value in row:
            if value is None or isinstance(value, bool):
                continue
            if isinstance(value, str):
                text(value, "单元格", 1000, empty=True)
            elif isinstance(value, (int, float)):
                if not math.isfinite(value) or not -1e12 < value < 1e12:
                    raise ValueError("单元格数字必须有限且绝对值小于一万亿")
            else:
                raise ValueError("单元格仅支持文本、数字、布尔或空值")
    return columns, rows


def write_table(output, basename, title, columns, rows, kind):
    """将已校验表格序列化，表格数据不经过公式求值或 HTML 执行。"""
    path = output / f"{basename}.{kind}"
    if kind == "csv":
        # UTF-8 BOM 便于导师用本地 Excel 打开中文。
        with path.open("w", encoding="utf-8-sig", newline="") as stream:
            writer = csv.writer(stream)
            writer.writerow([spreadsheet_literal(value) for value in columns])
            writer.writerows([[spreadsheet_literal(value) for value in row] for row in rows])
    elif kind == "xlsx":
        if any(isinstance(value, str) and len(value) > 32_767 for row in [columns, *rows] for value in row):
            raise ValueError("Excel 单元格最多 32767 字符，请改用其他格式或拆分内容")
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        from openpyxl.utils import get_column_letter

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "数据"
        sheet.append([spreadsheet_literal(value) for value in columns])
        for row in rows:
            sheet.append([spreadsheet_literal(value) for value in row])
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="3156A0")
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for index in range(1, len(columns) + 1):
            sheet.column_dimensions[get_column_letter(index)].width = 22
        workbook.save(path)
        workbook.close()
    elif kind == "json":
        path.write_text(json.dumps({"title": title, "columns": columns, "rows": rows},
                                   ensure_ascii=False, indent=2, allow_nan=False), encoding="utf-8")
    elif kind == "md":
        def cell(value):
            """保留普通单元格含义，同时阻止新行破坏表格结构。"""
            return markdown_literal("" if value is None else str(value)).replace("\r", " ").replace("\n", " / ")

        lines = [f"# {markdown_literal(title)}", "", "| " + " | ".join(cell(value) for value in columns) + " |",
                 "| " + " | ".join("---" for _ in columns) + " |"]
        lines.extend("| " + " | ".join(cell(value) for value in row) + " |" for row in rows)
        path.write_text("\n".join(lines) + "\n", encoding="utf-8")
    else:
        raise ValueError("表格格式仅支持 csv/xlsx/json/md")


def render_table(payload, output):
    """输出可下载的 CSV、Excel、JSON 或 Markdown 表格。"""
    title = text(payload.get("title"), "表格标题", 160)
    kind = format_name(payload.get("format", "xlsx"))
    if kind not in {"csv", "xlsx", "json", "md"}:
        raise ValueError("表格格式仅支持 csv/xlsx/json/md")
    columns, rows = validate_table(payload.get("columns"), payload.get("rows"))
    write_table(output, "table", title, columns, rows, kind)
    return f"表格已生成（{kind}）"


def render_document(payload, output):
    """从纯文本标题、正文和章节生成文档，不接受任意 HTML 模板。"""
    title = text(payload.get("title"), "文档标题", 160)
    content = text(payload.get("content", ""), "文档正文", 50_000, empty=True)
    kind = format_name(payload.get("format", "md"))
    raw_sections = payload.get("sections", [])
    if not isinstance(raw_sections, list) or len(raw_sections) > 12:
        raise ValueError("文档最多 12 个章节")
    sections = []
    for section in raw_sections:
        if not isinstance(section, dict):
            raise ValueError("章节需要 heading 和 body")
        sections.append({"heading": text(section.get("heading"), "章节标题", 100),
                         "body": text(section.get("body"), "章节正文", 8000, empty=True)})
    if kind in {"csv", "xlsx"}:
        rows = [["标题", title]] + ([["正文", content]] if content else [])
        rows += [[section["heading"], section["body"]] for section in sections]
        # 文档章节可长于普通表格单元格；格式范围仍受入口和 4 MB 产物限制。
        write_table(output, "document", title, ["章节", "正文"], rows, kind)
        return f"文档已生成（{kind}）"
    if kind == "json":
        value = json.dumps({"title": title, "content": content, "sections": sections},
                           ensure_ascii=False, indent=2, allow_nan=False)
    elif kind == "txt":
        value = title + "\n\n" + content
        value += "".join("\n\n" + section["heading"] + "\n" + section["body"] for section in sections)
    elif kind == "md":
        value = f"# {markdown_literal(title)}\n\n{markdown_literal(content)}"
        value += "".join(f"\n\n## {markdown_literal(section['heading'])}\n\n{markdown_literal(section['body'])}"
                         for section in sections)
    elif kind == "html":
        # 即使下载后打开，CSP 仍禁止脚本、网络资源和表单提交；所有正文均按文本转义。
        value = ('<!doctype html><html lang="zh-CN"><head><meta charset="utf-8">'
                 '<meta http-equiv="Content-Security-Policy" content="default-src \'none\'; '
                 'base-uri \'none\'; form-action \'none\'">'
                 f'<title>{html.escape(title)}</title></head><body><h1>{html.escape(title)}</h1>'
                 f'<pre>{html.escape(content)}</pre>')
        value += "".join(f"<h2>{html.escape(section['heading'])}</h2><pre>{html.escape(section['body'])}</pre>"
                         for section in sections)
        value += "</body></html>"
    (output / f"document.{kind}").write_text(value + "\n", encoding="utf-8")
    return f"文档已生成（{kind}）"
