"""一次性分析容器入口。仅通过 stdin 接收有限 JSON，不接触宿主目录或凭证。"""

import base64
import contextlib
import io
import json
from pathlib import Path
import sys

OUTPUT = Path("/tmp/output")
INPUT = Path("/tmp/input")
MAX_ARTIFACT = 4_000_000
ALLOWED_SUFFIXES = {".png", ".pdf", ".docx", ".csv", ".txt", ".xlsx", ".md", ".html", ".json"}


def chart(payload):
    """保留原入口，将具体绘图和维度校验交给独立渲染模块。"""
    from renderers.charts import render_chart
    return render_chart(payload, OUTPUT)


def report(payload):
    """保留原有 Word 报告入口，不把模型内容作为模板代码运行。"""
    from docx import Document
    from docx.shared import Pt

    sections = payload["sections"]
    if not 1 <= len(sections) <= 12:
        raise ValueError("报告需要 1–12 个章节")
    doc = Document()
    doc.add_heading(payload["title"], 0)
    for section in sections:
        doc.add_heading(str(section["heading"])[:100], level=1)
        doc.add_paragraph(str(section["body"])[:5000])
    for style_name in ("Normal", "Title"):
        doc.styles[style_name].font.name = "Microsoft YaHei"
        doc.styles[style_name].font.size = Pt(11 if style_name == "Normal" else 20)
    doc.save(OUTPUT / "report.docx")
    return "Word 报告已生成"


def execute_code(payload):
    """在容器内运行分析代码；输入和输出分开，保留代码主动生成的 data.json。"""
    # Docker 是隔离边界；此处明确运行模型提供的 Python，而不假装 AST 黑名单能构成沙箱。
    code = payload["code"]
    if len(code) > 6000:
        raise ValueError("分析代码最多 6000 字符")
    data = payload["data"]
    INPUT.mkdir(parents=True, exist_ok=True)
    input_file = INPUT / "data.json"
    input_file.write_text(json.dumps(data, ensure_ascii=False), encoding="utf-8")
    scope = {"__name__": "__main__", "data": data, "OUTPUT_DIR": str(OUTPUT), "INPUT_FILE": str(input_file)}
    stdout = io.StringIO()
    with contextlib.redirect_stdout(stdout):
        exec(compile(code, "<analysis>", "exec"), scope)
    return stdout.getvalue()[:12000]


def main():
    """分派有限输入任务，只收集允许格式且大小合规的普通输出文件。"""
    OUTPUT.mkdir(parents=True, exist_ok=True)
    raw = sys.stdin.buffer.read(100_001)
    if len(raw) > 100_000:
        raise ValueError("输入数据超过 100 KB")
    payload = json.loads(raw)
    mode = payload["mode"]
    if mode == "chart":
        message = chart(payload)
    elif mode == "report":
        message = report(payload)
    elif mode == "code":
        message = execute_code(payload)
    elif mode == "document":
        from renderers.documents import render_document
        message = render_document(payload, OUTPUT)
    elif mode == "table":
        from renderers.documents import render_table
        message = render_table(payload, OUTPUT)
    else:
        raise ValueError("未知分析任务")
    artifacts = []
    for path in sorted(OUTPUT.iterdir()):
        if path.suffix.lower() not in ALLOWED_SUFFIXES:
            continue
        if path.is_symlink() or not path.is_file():
            raise ValueError("输出必须是普通文件")
        if not 0 < path.stat().st_size <= MAX_ARTIFACT:
            raise ValueError("生成文件为空或超过 4 MB")
        artifacts.append({"name": path.name,
                          "content": base64.b64encode(path.read_bytes()).decode("ascii")})
        if len(artifacts) > 5:
            raise ValueError("最多生成 5 个文件")
    print(json.dumps({"ok": True, "message": message, "artifacts": artifacts}, ensure_ascii=False))


if __name__ == "__main__":
    try:
        main()
    except Exception as exc:
        # 不回传内部路径或堆栈，避免模型把它们当作业务事实。
        print(json.dumps({"ok": False, "error": str(exc)[:300]}, ensure_ascii=False))
